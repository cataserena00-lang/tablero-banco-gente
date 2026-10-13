"""Busca y baja las bases semanales de la CARPETA de Drive donde Desarrollo Social las carga.

Cada semana se mueven los Excel de la semana anterior a la subcarpeta «Versiones anteriores» y se cargan, con estos nombres,
las bases nuevas:
  - BASE PENDIENTES            créditos aprobados pendientes de entrega
  - BASE ESTADO DE CREDITOS    todas las solicitudes y su estado

El nombre se compara sin extensión, sin tildes y sin distinguir mayúsculas, y solo en la carpeta (no en las subcarpetas). Si hay
más de un archivo con ese nombre se usa el modificado más recientemente.

Variables de entorno:
  BG_DRIVE_CARPETA_ID            ID de la carpeta (el tramo largo de la URL). Si no está, se usan los IDs de archivo de antes
                                 (BG_DRIVE_FILE_ID para pendientes, BG_ESTADOS_FILE_ID para estados).
  GOOGLE_SERVICE_ACCOUNT_JSON    cuenta de servicio con acceso de lectura a la carpeta.

Estos archivos traen datos NOMINALES: este módulo no imprime filas ni valores, solo nombres de archivo y cantidades.
"""
import csv
import io
import os
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

NOMBRE_PENDIENTES = "BASE PENDIENTES"
NOMBRE_ESTADOS = "BASE ESTADO DE CREDITOS"
CARPETA = "application/vnd.google-apps.folder"
HOJA_GOOGLE = "application/vnd.google-apps.spreadsheet"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def normalizar_nombre(nombre: str) -> str:
    """'Base Estado de Créditos.xlsx' -> 'BASE ESTADO DE CREDITOS' (sin extensión, sin tildes, mayúsculas, espacios simples)."""
    base = re.sub(r"\.(xlsx|xls|csv)$", "", nombre.strip(), flags=re.I)
    sin = "".join(c for c in unicodedata.normalize("NFKD", base) if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", sin).strip().upper()


def elegir(archivos: list[dict], nombre: str) -> dict | None:
    """El archivo (no carpeta) de la lista con ese nombre; si hay varios, el modificado más recientemente."""
    objetivo = normalizar_nombre(nombre)
    cands = [a for a in archivos if a.get("mimeType") != CARPETA and normalizar_nombre(a.get("name", "")) == objetivo]
    return max(cands, key=lambda a: a.get("modifiedTime", ""), default=None)


def servicio():
    import json
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    info = json.loads(os.environ["GOOGLE_SERVICE_ACCOUNT_JSON"])
    cred = service_account.Credentials.from_service_account_info(info, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build("drive", "v3", credentials=cred, cache_discovery=False)


def listar(svc, carpeta_id: str) -> list[dict]:
    archivos, token = [], None
    while True:
        r = svc.files().list(
            q=f"'{carpeta_id}' in parents and trashed = false", pageSize=100, pageToken=token,
            fields="nextPageToken, files(id, name, mimeType, modifiedTime, size)",
            supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
        archivos += r.get("files", [])
        token = r.get("nextPageToken")
        if not token:
            return archivos


def bajar(svc, archivo: dict, destino: Path) -> None:
    from googleapiclient.http import MediaIoBaseDownload
    if archivo["mimeType"] == HOJA_GOOGLE:   # hoja nativa de Google: se exporta a xlsx (Google limita el tamaño de la exportación)
        pedido = svc.files().export_media(fileId=archivo["id"], mimeType=XLSX)
    else:
        pedido = svc.files().get_media(fileId=archivo["id"], supportsAllDrives=True)
    with open(destino, "wb") as f:
        dl = MediaIoBaseDownload(f, pedido, chunksize=16 * 1024 * 1024)
        hecho = False
        while not hecho:
            _, hecho = dl.next_chunk()


def bajar_por_nombre(carpeta_id: str, nombre: str, destino: Path) -> dict:
    """Baja a `destino` la base con ese nombre de la carpeta. Si no está, falla diciendo qué archivos hay (solo nombres)."""
    svc = servicio()
    archivos = listar(svc, carpeta_id)
    elegido = elegir(archivos, nombre)
    if not elegido:
        hay = ", ".join(sorted(a["name"] for a in archivos if a["mimeType"] != CARPETA)) or "ningún archivo"
        raise SystemExit(f"No se encontró «{nombre}» en la carpeta de Drive (o la cuenta de servicio no tiene acceso). Archivos visibles: {hay}.")
    bajar(svc, elegido, destino)
    print(f"Base «{elegido['name']}» (modificada {elegido['modifiedTime']}) descargada de la carpeta.")
    return elegido


def es_xlsx(ruta: Path) -> bool:
    """Los .xlsx son zips: empiezan con «PK»."""
    with open(ruta, "rb") as f:
        return f.read(2) == b"PK"


# ── xlsx -> csv (en streaming: el Excel de estados pesa decenas de MB y tiene cientos de miles de filas) ──────────────
PALABRAS_MONTO = re.compile(r"monto|deuda|valor|recupero|cuota|importe|prestable", re.I)


def texto_celda(v, encabezado: str) -> str:
    """Celda como texto, parecido a como la traía el CSV de antes: fechas como d/m/aaaa, enteros sin «.0» y los montos con
    puntos de miles (300.000). Vacío si no hay valor."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return f"{v.day}/{v.month}/{v.year}" if (v.hour, v.minute, v.second) == (0, 0, 0) else f"{v.day}/{v.month}/{v.year} {v:%H:%M:%S}"
    if isinstance(v, date):
        return f"{v.day}/{v.month}/{v.year}"
    if isinstance(v, bool):
        return "SI" if v else "NO"
    if isinstance(v, float):
        if v.is_integer():
            v = int(v)
        else:
            return str(v).replace(".", ",") if PALABRAS_MONTO.search(encabezado) else str(v)
    if isinstance(v, int):
        return f"{v:,}".replace(",", ".") if PALABRAS_MONTO.search(encabezado) and abs(v) >= 1000 else str(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def xlsx_a_csv(origen: Path, destino: Path) -> int:
    """Convierte la primera hoja del xlsx a CSV (utf-8, «;»). Devuelve la cantidad de filas de datos. No imprime valores."""
    from openpyxl import load_workbook
    wb = load_workbook(origen, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    n = 0
    with open(destino, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        filas = ws.iter_rows(values_only=True)
        enc = next(filas, None)
        if enc is None:
            raise SystemExit("La primera hoja del Excel está vacía.")
        nombres = ["" if x is None else re.sub(r"\s+", " ", str(x)).strip() for x in enc]
        while nombres and nombres[-1] == "":   # columnas vacías al final
            nombres.pop()
        w.writerow(nombres)
        for fila in filas:
            celdas = [texto_celda(v, nombres[i]) if i < len(nombres) else "" for i, v in enumerate(fila[:len(nombres)])]
            if any(celdas):
                w.writerow(celdas + [""] * (len(nombres) - len(celdas)))
                n += 1
    wb.close()
    return n


def estados_a_csv(carpeta_id: str, destino_csv: Path) -> None:
    """Baja «BASE ESTADO DE CREDITOS» de la carpeta y la deja como CSV en `destino_csv` (convirtiéndola si es un Excel)."""
    crudo = destino_csv.with_name(destino_csv.stem + "_descarga.xlsx")   # openpyxl solo abre archivos con extensión de Excel
    bajar_por_nombre(carpeta_id, NOMBRE_ESTADOS, crudo)
    if es_xlsx(crudo):
        filas = xlsx_a_csv(crudo, destino_csv)
        crudo.unlink()
        print(f"Excel convertido a CSV: {filas} filas.")
    else:
        crudo.replace(destino_csv)


def pendientes_bytes(carpeta_id: str) -> bytes:
    """Baja «BASE PENDIENTES» de la carpeta y devuelve el xlsx."""
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "pendientes.xlsx"
        bajar_por_nombre(carpeta_id, NOMBRE_PENDIENTES, ruta)
        return ruta.read_bytes()
