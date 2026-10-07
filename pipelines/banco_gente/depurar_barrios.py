"""
Depuración de barrios de la ciudad de Córdoba.

Para cada persona con Localidad = CORDOBA toma el texto que está después del
último "-" del Domicilio y lo asigna a un barrio oficial de barrios_cordoba.xlsx.

Uso:
    python depurar_barrios.py
    python depurar_barrios.py "base.xlsx" "barrios_cordoba.xlsx" "salida.xlsx"

Columnas que agrega a la base (al final):
    Barrio_Original   texto crudo que venía después del "-"
    Barrio_Depurado   barrio oficial asignado (vacío si no hay dato o no se pudo asignar)
    Barrio_Metodo     cómo se resolvió (EXACTO, NORMALIZADO, ALIAS, FONETICO, ...)
    Barrio_Candidatos si quedó ambiguo / sin match: opciones o texto a revisar
Además crea la hoja "Revision_Barrios" con los valores únicos que no se resolvieron.
"""
import re
import sys
import unicodedata
import warnings
from collections import Counter

import openpyxl
from openpyxl.styles import Font
from rapidfuzz import fuzz, process

warnings.filterwarnings("ignore")

BASE = r"C:\Users\catas\Downloads\base pendientes 6-10 (1).xlsx"
BARRIOS = r"C:\Users\catas\Downloads\barrios_cordoba.xlsx"
SALIDA = "base_pendientes_con_barrio.xlsx"

COL_LOCALIDAD = "Localidad"
COL_DOMICILIO = "Domicilio"
LOCALIDAD_OBJETIVO = "CORDOBA"

UMBRAL_FUZZY = 90      # similitud mínima (0-100) para aceptar un match aproximado
MARGEN_FUZZY = 4       # el mejor candidato debe superar al segundo por este margen

# ---------------------------------------------------------------------------
# Reglas de normalización
# ---------------------------------------------------------------------------
# abreviaturas / variantes -> forma canónica (se aplican token a token)
SINONIMOS = {
    "GRAL": "GENERAL", "GRL": "GENERAL",
    "AMPL": "AMPLIACION", "AMP": "AMPLIACION", "AMPLIAC": "AMPLIACION",
    "ANEXA": "ANEXO", "ANEX": "ANEXO", "ANX": "ANEXO",
    "TTE": "TENIENTE", "CNEL": "CORONEL", "CRL": "CORONEL", "CMTE": "COMANDANTE", "ALTE": "ALMIRANTE",
    "DR": "DOCTOR", "DRA": "DOCTOR", "STA": "SANTA", "STO": "SANTO",
    "PQUE": "PARQUE", "PQ": "PARQUE", "ALTOS": "ALTO",
    "RES": "RESIDENCIAL", "RESID": "RESIDENCIAL",
    "SECC": "SECCION", "SEC": "SECCION", "SECCIONES": "SECCION",
    "COOP": "COOPERATIVA", "URB": "URBANIZACION",
    "BARRIO": "", "BO": "", "B": "",          # "B° Alberdi", "Bº Alberdi"
    # puntos cardinales
    "SUR": "SUD", "N": "NORTE", "S": "SUD", "E": "ESTE", "O": "OESTE",
    "NOR": "NORTE", "OSTE": "OESTE",
    # ordinales -> número
    "PRIMERA": "1", "PRIMERO": "1", "1RA": "1", "1ERA": "1", "1A": "1", "1ER": "1", "1RO": "1", "UNO": "1",
    "SEGUNDA": "2", "SEGUNDO": "2", "2DA": "2", "2A": "2", "2DO": "2", "II": "2", "DOS": "2",
    "TERCERA": "3", "TERCERO": "3", "3RA": "3", "3ERA": "3", "3A": "3", "3RO": "3", "III": "3", "TRES": "3",
    "CUARTA": "4", "CUARTO": "4", "4TA": "4", "4A": "4", "IV": "4",
    "QUINTA": "5", "5TA": "5",
}
# palabras que no distinguen un barrio de otro
STOPWORDS = {"DE", "DEL", "LA", "LAS", "LOS", "EL", "Y", "A", "SECCION"}
# "SECCION" se descarta pero el número queda: "SANTA ISABEL 2DA SECCION" -> SANTA ISABEL 2

# Palabras que sólo agregan ruido al nombre del barrio (se sacan si el texto no coincide tal cual)
RUIDO = {"IPV", "ASENTAMIENTO", "URBANIZACION"}
# Barrios oficiales de una sola palabra demasiado genéricos para usarlos como "texto contenido"
GENERICOS = {"JARDIN", "AVENIDA", "CENTRO", "COLON", "MIRADOR", "EMPALME", "PATRIA", "COMERCIAL",
             "POLICIAL", "OBRERO", "INDUSTRIAL", "CONGRESO", "CALIFORNIA", "AEROPUERTO", "HORIZONTE"}

# Palabras que pueden sobrar o faltar sin cambiar de barrio (se prueban en un paso aparte)
DESCRIPTORES = {"VILLA", "GENERAL", "PARQUE", "RESIDENCIAL", "BARRIO", "CIUDAD", "JARDIN",
                "DOCTOR", "TENIENTE", "CORONEL", "COMANDANTE", "SANTA", "SAN", "ALTO", "BAJO", "ALTOS"}
# Calificativos que, si el barrio exacto no existe, se pueden quitar para caer en el barrio base
CALIFICATIVOS = {"AMPLIACION", "ANEXO", "NORTE", "SUD", "ESTE", "OESTE", "1", "2", "3", "4"}

# Equivalencias manuales de casos que no resuelve ninguna regla general.
# clave: texto ya normalizado (ver normalizar()) -> nombre OFICIAL (tal cual figura en barrios_cordoba.xlsx)
ALIAS = {
    "PUEYRREDON": "GENERAL PUEYRREDON",
    "CABILDO": "EL CABILDO",
    "MATIENZO": "TTE. BENJAMIN MATIENZO",
    "ZUMARAN": "ANA MARIA ZUMARAN",
    "ARENALES": "GENERAL ARENALES",
    "DON BOSCO": "PARQUE DON BOSCO",
    "JORGE NEWBERY": "PARQUE JORGE NEWBERY",
    "CORONEL OLMEDO": "VILLA CORONEL OLMEDO",
    "VICOR": "V.I.C.O.R.",
    "SEP": "S.E.P.",
    "UOCRA": "U.O.C.R.A.",
    "OBRAS SANITARIAS": "O.S.N.",
    "OSN": "O.S.N.",
    "ATE": "A.T. E.",
    "COVICO": "CO.VI.CO.",
    "VILLA URQUIZA": "VILLA GENERAL URQUIZA",     # también existe URQUIZA: revisar si hace falta distinguir
    "GENERAL URQUIZA": "VILLA GENERAL URQUIZA",
    "AMPLIACION AMERICA": "AMPLIACION RESIDENCIAL AMERICA",
    "CARCANO": "RAMON J. CARCANO",
    "ALEJANDRO CARBO": "CARBO",
    "AVELLANEDA": "NICOLAS AVELLANEDA",
    "REMO COPELLO": "DR REMO M. COPELLO",
}

# valores que significan "no hay barrio"
SIN_DATO = {"", "SD", "S D", "S/D", "SIN DATO", "SIN DATOS", "NO", "NO TIENE", "NC", "N C", "0", "X", "XX"}


def sin_tildes(t):
    """Quita tildes y diéresis pero conserva la Ñ."""
    t = t.replace("Ñ", "\x01").replace("ñ", "\x02")
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    return t.replace("\x01", "Ñ").replace("\x02", "ñ")


def normalizar(texto):
    """Texto -> clave canónica: mayúsculas, sin tildes ni puntuación, abreviaturas expandidas,
    ordinales a número, SUR=SUD, sin palabras de relleno."""
    t = sin_tildes(str(texto)).upper()
    t = t.replace("º", "").replace("°", "").replace("ª", "")
    t = re.sub(r"[^A-Z0-9Ñ ]", " ", t)                  # puntos, guiones, comillas, paréntesis...
    t = unir_siglas(re.sub(r"\s+", " ", t).strip())     # "V.I.C.O.R." -> "VICOR", "O.S.N." -> "OSN"
    t = re.sub(r"\bJOSE I DIAZ\b|\bJ I DIAZ\b", "JOSE IGNACIO DIAZ", t)
    crudos = t.split()
    if len(crudos) > 1 and crudos[-1] == "I":            # "LAS FLORES I" -> "LAS FLORES 1"
        crudos[-1] = "1"
    toks = []
    for tok in crudos:
        tok = SINONIMOS.get(tok, tok)
        if tok and tok not in STOPWORDS:
            toks.append(tok)
    return " ".join(toks)


def unir_siglas(t):
    """Letras sueltas consecutivas son una sigla: 'V I C O R' -> 'VICOR'."""
    return re.sub(r"\b(?:[A-Z] )+[A-Z]\b", lambda m: m.group(0).replace(" ", ""), t)


def fonetica(clave):
    """Clave tolerante a errores de ortografía (S/Z/C, V/B, Y/I/LL, H muda, letras dobles)."""
    out = []
    for tok in clave.split():
        if tok.isdigit():
            out.append(tok)
            continue
        t = tok
        t = re.sub(r"C(?=[EI])", "S", t)
        t = t.replace("QU", "K").replace("C", "K").replace("Z", "S")
        t = t.replace("LL", "I").replace("Y", "I").replace("V", "B").replace("W", "U")
        t = t.replace("H", "").replace("Ñ", "N")
        t = re.sub(r"(.)\1+", r"\1", t)
        out.append(t)
    return " ".join(out)


def clave_total(texto):
    return normalizar(texto)


# ---------------------------------------------------------------------------
# Extraer el barrio del domicilio
# ---------------------------------------------------------------------------
def extraer_barrio(domicilio):
    """Devuelve el texto posterior al último ' - ' (o al último '-' si no hay con espacios)."""
    if domicilio is None:
        return ""
    d = str(domicilio)
    if " - " in d:
        return d.rpartition(" - ")[2].strip()
    if "-" in d:
        return d.rpartition("-")[2].strip()
    return ""   # sin guion -> no hay campo de barrio


# ---------------------------------------------------------------------------
# Matcher
# ---------------------------------------------------------------------------
class MatcherBarrios:
    def __init__(self, ruta_barrios):
        wb = openpyxl.load_workbook(ruta_barrios, read_only=True, data_only=True)
        ws = wb["Barrios"]
        filas = list(ws.iter_rows(values_only=True))
        h = [str(x).strip() for x in filas[0]]
        i = h.index("Barrio")
        oficiales = []
        for f in filas[1:]:
            b = f[i]
            if b is None:
                continue
            b = str(b).strip()
            if b and b.upper() != "SD" and b not in oficiales:   # "SD" = sin dato; hay duplicados
                oficiales.append(b)
        self.oficiales = oficiales
        self.por_clave = {}        # clave normalizada -> [oficiales]
        for b in oficiales:
            self.por_clave.setdefault(clave_total(b), []).append(b)
        self.por_fon = {}          # clave fonética -> [oficiales]
        self.por_set = {}          # palabras ordenadas -> [oficiales]
        self.por_pegado = {}       # sin espacios ("LA SALLE" = "LASALLE")
        for k, v in self.por_clave.items():
            self.por_fon.setdefault(fonetica(k), []).extend(v)
            self.por_set.setdefault(" ".join(sorted(k.split())), []).extend(v)
        for b in oficiales:
            self.por_pegado.setdefault(self._pegado(b), []).append(b)
        self.alias = {clave_total(a): b for a, b in ALIAS.items()}
        for b in self.alias.values():
            if b not in oficiales:
                raise ValueError(f"El alias apunta a un barrio que no existe en la lista oficial: {b}")
        self.cache = {}

    # --- utilidades -------------------------------------------------------
    @staticmethod
    def _unico(lista):
        lista = list(dict.fromkeys(lista))
        return lista[0] if len(lista) == 1 else None

    @staticmethod
    def _pegado(texto):
        t = sin_tildes(str(texto)).upper().replace("SUR", "SUD")
        return fonetica(re.sub(r"[^A-Z0-9Ñ]", "", t))

    @staticmethod
    def _sin_calificativos(clave):
        return " ".join(t for t in clave.split() if t not in CALIFICATIVOS)

    @staticmethod
    def _variantes_descriptor(clave):
        """Claves que resultan de agregar o quitar UNA palabra descriptora (VILLA, GENERAL, PARQUE...)."""
        toks = clave.split()
        vs = set()
        for i, t in enumerate(toks):
            if t in DESCRIPTORES and len(toks) > 1:
                vs.add(" ".join(toks[:i] + toks[i + 1:]))
        for d in DESCRIPTORES:
            vs.add(d + " " + clave)
            vs.add(clave + " " + d)
        return vs

    def _directo(self, clave):
        """Pasos directos. Devuelve (barrio, metodo) o None."""
        u = self._unico(self.por_clave.get(clave, []))
        if u:
            return u, "NORMALIZADO"
        if clave in self.alias:
            return self.alias[clave], "ALIAS"
        u = self._unico(self.por_set.get(" ".join(sorted(clave.split())), []))
        if u:
            return u, "ORDEN DE PALABRAS"
        u = self._unico(self.por_fon.get(fonetica(clave), []))
        if u:
            return u, "FONETICO"
        u = self._unico(self.por_pegado.get(self._pegado(clave), []))
        if u:
            return u, "FONETICO"
        return None

    def _descriptor(self, clave):
        """Prueba agregar/quitar VILLA, GENERAL, PARQUE, etc. Devuelve (barrio|None, candidatos)."""
        cands = []
        for v in self._variantes_descriptor(clave):
            for tabla, k in ((self.por_clave, v), (self.por_fon, fonetica(v))):
                cands += tabla.get(k, [])
        cands = list(dict.fromkeys(cands))
        return (cands[0] if len(cands) == 1 else None), cands

    # --- resolución -------------------------------------------------------
    def resolver(self, crudo):
        """Devuelve (barrio_oficial|'', metodo, candidatos)."""
        if crudo not in self.cache:
            self.cache[crudo] = self._resolver(crudo)
        return self.cache[crudo]

    def _resolver(self, crudo):
        if crudo is None or crudo.strip().upper() in SIN_DATO:
            return "", "SIN DATO", ""
        clave = clave_total(crudo)
        if not clave or clave in SIN_DATO:
            return "", "SIN DATO", ""

        # 1) idéntico al nombre oficial
        if crudo.strip() in self.oficiales:
            return crudo.strip(), "EXACTO", ""

        # 2) tildes, abreviaturas, SUR/SUD, ordinales, puntuación, orden, ortografía, alias
        r = self._directo(clave)
        if r:
            return r[0], r[1], ""

        # 3) sacar palabras de ruido (IPV, ASENTAMIENTO...) y reintentar
        toks = [t for t in clave.split() if t not in RUIDO]
        clave2 = " ".join(toks)
        if clave2 and clave2 != clave:
            r = self._directo(clave2)
            if r:
                return r[0], r[1] + " (sin ruido)", ""
            clave = clave2

        # 4) loteos de "Horizonte": "X de Horizonte" -> barrio oficial si existe, si no HORIZONTE
        if "HORIZONTE" in toks:
            r = self._directo(self._sin_calificativos(clave))
            if r:
                return r[0], "BARRIO BASE (sin ampliación/anexo/cardinal)", ""
            return "HORIZONTE", "LOTEO DE HORIZONTE", ""

        # 5) sobra/falta VILLA, GENERAL, PARQUE, RESIDENCIAL...
        u, cands = self._descriptor(clave)
        if u:
            return u, "DESCRIPTOR VILLA/GRAL/PARQUE", ""
        if len(cands) > 1:
            return "", "AMBIGUO", " | ".join(cands)

        # 6) falta la sección / número (PARQUE LICEO -> 1, 2 o 3)
        n = len(clave.split())
        prefijo = [b for k, bs in self.por_clave.items() for b in bs
                   if k.startswith(clave + " ") and k.split()[-1].isdigit() and len(k.split()) == n + 1]
        if len(prefijo) > 1:
            return "", "AMBIGUO", " | ".join(prefijo)

        # 7) varios números ("SECCIONES 1 2 3") -> no se puede elegir una
        if len({t for t in clave.split() if t.isdigit()}) > 1:
            base = self._sin_calificativos(clave)
            multi = [b for k, bs in self.por_clave.items() for b in bs if k.startswith(base + " ")]
            return "", "AMBIGUO", " | ".join(multi)

        # 8) quitar AMPLIACION / ANEXO / puntos cardinales / número y buscar el barrio base
        base = self._sin_calificativos(clave)
        if base and base != clave:
            r = self._directo(base)
            if r:
                return r[0], "BARRIO BASE (sin ampliación/anexo/cardinal)", ""
            u, _ = self._descriptor(base)
            if u:
                return u, "BARRIO BASE (sin ampliación/anexo/cardinal)", ""

        # 9) texto de más: el barrio oficial más largo contenido dentro del texto
        toks = clave.split()
        for largo in range(len(toks), 0, -1):
            hallados = []
            for i in range(len(toks) - largo + 1):
                sub = " ".join(toks[i:i + largo])
                if largo == 1 and (sub in GENERICOS or len(sub) < 5):
                    continue
                hallados += self.por_clave.get(sub, [])
            hallados = list(dict.fromkeys(hallados))
            if len(hallados) == 1:
                return hallados[0], "TEXTO SOBRANTE", ""
            if len(hallados) > 1:
                return "", "AMBIGUO", " | ".join(hallados)

        # 9b) el texto es sólo un pedazo del nombre oficial ("DOROTEA" -> "FINCA LA DOROTEA")
        if len(clave) >= 5 and clave not in GENERICOS:
            frag = [b for k, bs in self.por_clave.items() for b in bs
                    if f" {clave} " in f" {k} "]
            if len(frag) == 1:
                return frag[0], "FRAGMENTO DEL NOMBRE OFICIAL", ""
            if len(frag) > 1:
                return "", "AMBIGUO", " | ".join(frag)

        # 10) parecido (errores de tipeo); los números tienen que coincidir
        fon = fonetica(clave)
        nums = {t for t in clave.split() if t.isdigit()}
        claves = [k for k in self.por_fon if {t for t in k.split() if t.isdigit()} == nums]
        res = process.extract(fon, claves, scorer=fuzz.ratio, limit=3) if claves else []
        if res:
            mej, score = res[0][0], res[0][1]
            segundo = res[1][1] if len(res) > 1 else 0
            if score >= UMBRAL_FUZZY and score - segundo >= MARGEN_FUZZY:
                u = self._unico(self.por_fon[mej])
                if u:
                    return u, f"APROXIMADO ({score:.0f}%)", ""
            sug = " | ".join(f"{self._unico(self.por_fon[c]) or c} ({s:.0f}%)" for c, s, _ in res)
            return "", "SIN COINCIDENCIA", sug
        return "", "SIN COINCIDENCIA", ""


# ---------------------------------------------------------------------------
# Proceso principal
# ---------------------------------------------------------------------------
def main(base=BASE, barrios=BARRIOS, salida=SALIDA):
    m = MatcherBarrios(barrios)
    wb = openpyxl.load_workbook(base)
    ws = wb.active
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    c_loc, c_dom = headers[COL_LOCALIDAD], headers[COL_DOMICILIO]
    ult = ws.max_column
    nuevas = ["Barrio_Original", "Barrio_Depurado", "Barrio_Metodo", "Barrio_Candidatos"]
    for j, n in enumerate(nuevas, 1):
        c = ws.cell(row=1, column=ult + j, value=n)
        c.font = Font(bold=True)

    resumen, revision = Counter(), Counter()
    ejemplos = {}
    for r in range(2, ws.max_row + 1):
        loc = ws.cell(row=r, column=c_loc).value
        if loc is None or sin_tildes(str(loc)).strip().upper() != LOCALIDAD_OBJETIVO:
            continue
        crudo = extraer_barrio(ws.cell(row=r, column=c_dom).value)
        barrio, metodo, cands = m.resolver(crudo)
        for j, v in enumerate((crudo, barrio, metodo, cands), 1):
            ws.cell(row=r, column=ult + j, value=v if v != "" else None)
        resumen[metodo.split(" (")[0]] += 1
        if not barrio and metodo != "SIN DATO":
            revision[(crudo, metodo, cands)] += 1

    # hoja de revisión
    if "Revision_Barrios" in wb.sheetnames:
        del wb["Revision_Barrios"]
    wr = wb.create_sheet("Revision_Barrios")
    wr.append(["Texto original", "Motivo", "Sugerencias", "Cantidad"])
    for c in wr[1]:
        c.font = Font(bold=True)
    for (crudo, metodo, cands), n in sorted(revision.items(), key=lambda x: -x[1]):
        wr.append([crudo, metodo, cands, n])

    wb.save(salida)
    print(f"Archivo generado: {salida}\n")
    total = sum(resumen.values())
    print(f"Personas de {LOCALIDAD_OBJETIVO}: {total}")
    for k, v in resumen.most_common():
        print(f"  {k:<45}{v:>6}  {v / total:6.1%}")
    return m


if __name__ == "__main__":
    a = sys.argv[1:]
    main(*a) if a else main()
